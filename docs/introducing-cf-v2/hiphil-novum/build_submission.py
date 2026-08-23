#!/usr/bin/env python3
"""Build the anonymous HIPHIL Novum Word and PDF submission.

The official journal template supplies the page layout, headers, and named styles.
The LaTeX article remains the content source; this script only adapts presentation,
citations, visual assets, and submission-only front/back matter.
"""

from __future__ import annotations

import re
import shutil
import subprocess
import urllib.request
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pypandoc
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


HERE = Path(__file__).resolve().parent
ARTICLE_DIR = HERE.parent
SOURCE = ARTICLE_DIR / "intro-to-cf__rewrite.tex"
TEMPLATE = HERE / "template" / "hiphil-novum-template.docx"
ASSETS = HERE / "assets"
BUILD = HERE / "build"
DOCX = HERE / "intro-to-cf__hiphil-novum.docx"
PDF = HERE / "intro-to-cf__hiphil-novum.pdf"

TEMPLATE_URL = "https://tidsskrift.dk/hiphilnovum/libraryFiles/downloadPublic/295"
PANDOC = Path(pypandoc.get_pandoc_path())

CITATIONS = {
    "oosting2016": "Oosting 2016",
    "roorda2014laf": "Roorda et al. 2014",
    "ide2014laf": "Ide and Suderman 2014",
    "harris2020numpy": "Harris et al. 2020",
    "mckinney2010pandas": "McKinney 2010",
    "pedregosa2011scikit": "Pedregosa et al. 2011",
    "naaijer2016parallel": "Naaijer and Roorda 2016",
    "kalkman2015verbal": "Kalkman 2015",
    "kingham2016song": "Kingham 2016",
    "erwich2021psalms": "Erwich 2021",
    "hojgaard2021roles": "Højgaard 2021",
    "vanpeursen2019": "van Peursen 2019",
    "textfabricsoftware": "Roorda 2016",
    "openai2022chatgpt": "OpenAI 2022",
    "pytorchmmap": "PyTorch Contributors n.d.",
    "tensorflowmmap": "Google AI Edge n.d.",
    "numpymemmap": "NumPy Developers n.d.",
    "bhsa": "van Peursen, Sikkel, and Roorda 2015",
    "lxx": "Center for Biblical Languages and Computing 2024a",
    "n1904": "Center for Biblical Languages and Computing 2024b",
    "dss": "Jacobs, Naaijer, and Roorda 2017",
    "peshitta": "van Peursen et al. 2018",
    "syrnt": "Vlaardingerbroek and Roorda 2017",
    "sp": "Naaijer et al. 2020",
    "tischendorf": "Kingham 2018",
    "quran": "van Lit and Roorda 2017",
    "cuc": "Højgaard et al. 2020",
    "beyer2019reliable": "Beyer, Löwe, and Wendler 2019",
    "psutil": "Psutil Contributors n.d.",
    "modelcontextprotocol": "Model Context Protocol Contributors 2025",
}

FIGURES = {
    "fig_containment_hierarchy.pdf": "containment-hierarchy.png",
    "fig_memory_multicorpus.pdf": "memory-multicorpus.png",
    "fig_scaling_progressive.pdf": "scaling-progressive.png",
    "fig_latency_distribution.pdf": "latency-distribution.png",
}


def run(*args: str | Path, cwd: Path | None = None) -> None:
    subprocess.run([str(arg) for arg in args], cwd=cwd, check=True)


def ensure_template() -> None:
    TEMPLATE.parent.mkdir(parents=True, exist_ok=True)
    if not TEMPLATE.exists():
        urllib.request.urlretrieve(TEMPLATE_URL, TEMPLATE)


def render_assets() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    BUILD.mkdir(parents=True, exist_ok=True)
    for source_name, output_name in FIGURES.items():
        output = ASSETS / output_name
        run(
            "pdftocairo",
            "-png",
            "-singlefile",
            "-r",
            "300",
            ARTICLE_DIR / source_name,
            output.with_suffix(""),
        )

    source = SOURCE.read_text()
    tikz = re.search(
        r"(\\begin\{tikzpicture\}.*?\\end\{tikzpicture\})", source, re.DOTALL
    )
    if tikz is None:
        raise RuntimeError("MCP agent-experience diagram not found")
    picture = tikz.group(1)
    picture = picture.replace("0.82\\textwidth", "5.75in")
    picture = picture.replace("0.36\\textwidth", "2.48in")
    tex = rf"""\documentclass[tikz,border=8pt]{{standalone}}
\usepackage{{tikz}}
\usepackage{{cjhebrew}}
\usepackage[scaled=0.9]{{inconsolata}}
\usetikzlibrary{{arrows.meta,positioning}}
\definecolor{{codegray}}{{rgb}}{{0.5,0.5,0.5}}
\definecolor{{backcolour}}{{rgb}}{{0.95,0.95,0.92}}
\begin{{document}}
{picture}
\end{{document}}
"""
    diagram_tex = BUILD / "mcp-agent-experience.tex"
    diagram_tex.write_text(tex)
    run(
        "pdflatex",
        "-interaction=nonstopmode",
        "-halt-on-error",
        "-output-directory",
        BUILD,
        diagram_tex,
    )
    run(
        "pdftocairo",
        "-png",
        "-singlefile",
        "-r",
        "300",
        BUILD / "mcp-agent-experience.pdf",
        (ASSETS / "mcp-agent-experience").with_suffix(""),
    )


def replace_citations(text: str) -> str:
    def replace(match: re.Match[str]) -> str:
        keys = [key.strip() for key in match.group(1).split(",")]
        return "(" + "; ".join(CITATIONS[key] for key in keys) + ")"

    return re.sub(r"\\cite\{([^}]+)\}", replace, text)


def replace_listings(text: str) -> str:
    example_number = 0
    pattern = re.compile(
        r"\\begin\{lstlisting\}(?:\[(.*?)\])?\n(.*?)\\end\{lstlisting\}",
        re.DOTALL,
    )

    def replace(match: re.Match[str]) -> str:
        nonlocal example_number
        options = match.group(1) or ""
        code = match.group(2).rstrip()
        code = re.sub(r"\(\*@\\cjRL\{b\*:\}@\*\)", "בְּ", code)
        code = re.sub(r"\(\*@\\cjRL\{rE'\+siyt\}@\*\)", "רֵאשִׁית", code)
        code = re.sub(r"\(\*@\\cjRL\{b\*ArA'\}@\*\)", "בָּרָא", code)
        code = re.sub(r"\(\*@\\cjRL\{'E:loh\}@\*\)", "אֱלֹה", code)
        caption_match = re.search(
            r"caption\s*=\s*(?:\{([^{}]*(?:\\texttt\{[^{}]+\}[^{}]*)*)\}|([^,\n]+))",
            options,
        )
        caption = None
        if caption_match:
            caption = (caption_match.group(1) or caption_match.group(2)).strip()
            example_number += 1
        numbered = "numbers=none" not in options and caption is not None
        if numbered:
            lines = code.splitlines()
            width = len(str(len(lines)))
            code = "\n".join(
                f"{index:>{width}}  {line}" for index, line in enumerate(lines, 1)
            )
        result = f"\\begin{{verbatim}}\n{code}\n\\end{{verbatim}}"
        if caption:
            result = (
                f"\\begin{{center}}Example {example_number}. {caption}\\end{{center}}\n"
                + result
            )
        return result

    return pattern.sub(replace, text)


def prepare_latex() -> Path:
    text = SOURCE.read_text()
    text = text.replace("\\author{Cody Kingham}", "\\author{}")
    text = text.replace("\\date{Aug 2026}", "\\date{}")
    text = text.replace("\\begin{tabular}{lrrrrr}", "\\begin{tabular}{lrrrr}", 1)
    text = text.replace(
        "\\end{abstract}",
        "\\end{abstract}\n\n\\noindent \\textbf{Keywords:} corpus linguistics; Text-Fabric; memory mapping; AI agents",
        1,
    )
    # Two accidental repeated lines in the LaTeX source are suppressed only in this copy.
    text = text.replace(
        "\\subsection{Text-Fabric's Python Dictionaries}\n\\subsection{Text-Fabric's Python Dictionaries}",
        "\\subsection{Text-Fabric's Python Dictionaries}",
    )
    text = text.replace(
        "Context-Fabric exposes corpora for structured analysis. Corpora contain hierarchical\n"
        "Context-Fabric exposes corpora for structured analysis. Corpora contain hierarchical",
        "Context-Fabric exposes corpora for structured analysis. Corpora contain hierarchical",
    )
    text = text.replace(
        "along with any description metadata. \\texttt{get\\_text\\_formats}. The agent's interactions with these\n"
        "three is illustrated",
        "along with any description metadata. The agent's interactions with these\n"
        "three are illustrated",
    )
    example_references = {
        "lst:mcp_connection_instructions": 8,
        "lst:mcp_corpus_description_tools": 9,
        "lst:mcp_text_formats": 10,
        "lst:mcp_search_syntax_guide": 11,
    }
    for label, number in example_references.items():
        text = text.replace(f"Listing~\\ref{{{label}}}", f"Example {number}")
    text = replace_citations(text)
    text = replace_listings(text)
    text = text.replace("\\cjRL{rE'+siyt}", "רֵאשִׁית")
    for source_name, output_name in FIGURES.items():
        text = text.replace(source_name, str((ASSETS / output_name).resolve()))

    diagram = re.compile(
        r"\\begin\{figure\}\[H\].*?\\end\{figure\}", re.DOTALL
    )
    replacement = rf"""\begin{{figure}}[htbp]
\centering
\includegraphics[width=0.94\textwidth]{{{(ASSETS / 'mcp-agent-experience.png').resolve()}}}
\caption{{The agent's progressive view through the MCP server.}}
\label{{fig:mcp_agent_experience}}
\end{{figure}}"""
    text, count = diagram.subn(lambda _: replacement, text, count=1)
    if count != 1:
        raise RuntimeError("MCP diagram figure was not replaced")

    figure_number = 0
    table_number = 0

    def number_figure(match: re.Match[str]) -> str:
        nonlocal figure_number
        figure_number += 1
        return match.group(0).replace(
            "\\caption{", f"\\caption{{Figure {figure_number}. ", 1
        )

    def number_table(match: re.Match[str]) -> str:
        nonlocal table_number
        table_number += 1
        return match.group(0).replace(
            "\\caption{", f"\\caption{{Table {table_number}. ", 1
        )

    text = re.sub(
        r"\\begin\{figure\}.*?\\end\{figure\}", number_figure, text, flags=re.DOTALL
    )
    text = re.sub(
        r"\\begin\{table\}.*?\\end\{table\}", number_table, text, flags=re.DOTALL
    )

    text = text.replace(
        "readers may follow the CF GitHub organization, at \\url{https://github.com/Context-Fabric/},\n"
        "or visit the CF website at \\url{https://context-fabric.ai}.",
        "readers may follow the CF GitHub organization,\\footnote{\\url{https://github.com/Context-Fabric/}}\n"
        "or visit the CF website.\\footnote{\\url{https://context-fabric.ai}}",
    )
    text = text.replace(
        "\\begin{thebibliography}{99}",
        "\\section*{Acknowledgements}\n\n"
        "I would like to acknowledge the scholarship, influence, and friendship of Dirk Roorda, "
        "whose creativity and rebelliousness have always inspired me. "
        "May his contributions continue to bear fruit into the distant future.\n\n"
        "During the preparation of this manuscript, the author used Anthropic Claude Opus 4.5 "
        "for the analysis code, and OpenAI Codex (GPT-5.6 Sol) for reference verification, "
        "preparation of visual elements, and document formatting. "
        "The re-write of TF -> CF was done entirely by Opus 4.5. "
        "The entirety of the body text was hand-written by the author, but based loosely on a draft "
        "text generated by Claude. "
        "The author reviewed and edited the output and takes full responsibility for the content "
        "of this publication.\n\n"
        "\\section*{Bibliography}",
    )
    text = re.sub(r"\\bibitem\{[^}]+\}\n", "", text)
    text = text.replace(
        "Center for Biblical Languages and Computing. (2024). Septuagint",
        "Center for Biblical Languages and Computing. (2024a). Septuagint",
    )
    text = text.replace(
        "Center for Biblical Languages and Computing. (2024). Nestle",
        "Center for Biblical Languages and Computing. (2024b). Nestle",
    )
    text = text.replace("\\end{thebibliography}", "")
    output = BUILD / "article-for-word.tex"
    output.write_text(text)
    return output


def set_border(parent, edge: str, value: str, size: int = 0) -> None:
    """Set one Word table or cell border."""
    border = parent.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        parent.append(border)
    border.set(qn("w:val"), value)
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), "000000")


def format_booktabs_table(table) -> None:
    """Apply a LaTeX booktabs-like rule and padding treatment."""
    table_properties = table._tbl.tblPr
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    set_border(borders, "top", "single", 10)
    set_border(borders, "bottom", "single", 10)
    for edge in ("left", "right", "insideH", "insideV"):
        set_border(borders, edge, "nil")

    for row_index, row in enumerate(table.rows):
        row_properties = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        row_properties.append(cant_split)
        for cell in row.cells:
            cell_properties = cell._tc.get_or_add_tcPr()
            shading = cell_properties.find(qn("w:shd"))
            if shading is not None:
                cell_properties.remove(shading)

            cell_margins = cell_properties.find(qn("w:tcMar"))
            if cell_margins is None:
                cell_margins = OxmlElement("w:tcMar")
                cell_properties.append(cell_margins)
            for edge, width in (("top", 55), ("bottom", 55), ("left", 85), ("right", 85)):
                margin = cell_margins.find(qn(f"w:{edge}"))
                if margin is None:
                    margin = OxmlElement(f"w:{edge}")
                    cell_margins.append(margin)
                margin.set(qn("w:w"), str(width))
                margin.set(qn("w:type"), "dxa")

            cell_borders = cell_properties.find(qn("w:tcBorders"))
            if cell_borders is None:
                cell_borders = OxmlElement("w:tcBorders")
                cell_properties.append(cell_borders)
            for edge in ("top", "bottom", "left", "right", "insideH", "insideV"):
                set_border(cell_borders, edge, "nil")
            if row_index == 0:
                set_border(cell_borders, "top", "single", 10)
                set_border(cell_borders, "bottom", "single", 7)
            if row_index == len(table.rows) - 1:
                set_border(cell_borders, "bottom", "single", 10)


def set_run_font(run, name: str, size: float | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)


def style_document(input_docx: Path) -> None:
    document = Document(input_docx)
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    document.core_properties.title = "Introducing Context-Fabric"
    document.core_properties.subject = "Anonymous article submission"

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    try:
        inline_code_style = document.styles["Verbatim Char"]
    except KeyError:
        inline_code_style = document.styles.add_style(
            "Verbatim Char", WD_STYLE_TYPE.CHARACTER
        )
    inline_code_style.font.name = "Consolas"
    inline_code_style.font.size = Pt(10.5)
    inline_code_style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Consolas")

    if "HIPHIL Code" not in [style.name for style in document.styles]:
        code_style = document.styles.add_style("HIPHIL Code", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = document.styles["HIPHIL Code"]
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(10)
    code_style.paragraph_format.line_spacing = Pt(10)
    code_style.paragraph_format.left_indent = Cm(0.25)
    code_style.paragraph_format.right_indent = Cm(0.25)
    code_style.paragraph_format.keep_together = True

    if "HIPHIL Example title" not in [style.name for style in document.styles]:
        example_title_style = document.styles.add_style(
            "HIPHIL Example title", WD_STYLE_TYPE.PARAGRAPH
        )
    else:
        example_title_style = document.styles["HIPHIL Example title"]
    example_title_style.base_style = document.styles["HIPHIL Table title"]
    example_title_style.font.name = "Times New Roman"
    example_title_style.font.size = Pt(11)
    example_title_style.paragraph_format.space_before = Pt(16)
    example_title_style.paragraph_format.space_after = Pt(8)
    example_title_style.paragraph_format.keep_with_next = True

    bibliography_started = False
    bibliography_paragraphs = []
    def format_caption(paragraph, style_name: str, centered: bool) -> None:
        caption = paragraph.text
        label, separator, description = caption.partition(". ")
        paragraph.clear()
        label_run = paragraph.add_run(label + (". " if separator else ""))
        label_run.italic = True
        description_run = paragraph.add_run(description)
        description_run.italic = False
        paragraph.style = document.styles[style_name]
        if centered:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for index, paragraph in enumerate(document.paragraphs):
        value = paragraph.text.strip()
        original_style = paragraph.style.name
        if not value and not paragraph._p.xpath(".//w:drawing"):
            continue
        if index == 0:
            paragraph.style = document.styles["HIPHIL Title."]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(18)
        elif index == 1 and value == "Abstract":
            paragraph._element.getparent().remove(paragraph._element)
            continue
        elif index == 2:
            label = paragraph.add_run("Abstract: ")
            label.bold = True
            insertion_index = 1 if paragraph._p.pPr is not None else 0
            paragraph._p.remove(label._r)
            paragraph._p.insert(insertion_index, label._r)
            paragraph.style = document.styles["HIPHIL Abstract"]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(10)
        elif original_style.startswith("Heading 1") or value in {
            "Acknowledgements",
            "Bibliography",
        }:
            paragraph.style = document.styles["HIPHIL Heading1"]
            paragraph.paragraph_format.space_before = Pt(18)
            paragraph.paragraph_format.space_after = Pt(10)
            paragraph.paragraph_format.keep_with_next = True
            bibliography_started = value == "Bibliography"
        elif original_style.startswith("Heading 2"):
            paragraph.style = document.styles["HIPHIL Heading2"]
            paragraph.paragraph_format.space_before = Pt(14)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_with_next = True
        elif original_style.startswith("Heading 3"):
            paragraph.style = document.styles["HIPHIL Heading3"]
            paragraph.paragraph_format.space_before = Pt(12)
            paragraph.paragraph_format.space_after = Pt(6)
            paragraph.paragraph_format.keep_with_next = True
        elif value.startswith("Abstract:") or original_style == "Abstract":
            paragraph.style = document.styles["HIPHIL Abstract"]
        elif value.startswith("Keywords:"):
            paragraph.style = document.styles["HIPHIL Abstract"]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(18)
        elif original_style == "Source Code":
            paragraph.style = code_style
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_together = len(value.splitlines()) <= 32
            properties = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "F2F2EC")
            properties.append(shading)
            for run in paragraph.runs:
                set_run_font(run, "Consolas", 8)
        elif value.startswith("Example "):
            format_caption(paragraph, "HIPHIL Example title", centered=False)
            paragraph.paragraph_format.space_before = Pt(16)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_with_next = True
        elif value.startswith("Table "):
            format_caption(paragraph, "HIPHIL Table title", centered=False)
            paragraph.paragraph_format.space_before = Pt(16)
            paragraph.paragraph_format.space_after = Pt(8)
            paragraph.paragraph_format.keep_with_next = True
        elif value.startswith("Figure ") or original_style == "Caption":
            format_caption(paragraph, "HIPHIL Figure caption", centered=True)
            paragraph.paragraph_format.space_before = Pt(6)
            paragraph.paragraph_format.space_after = Pt(14)
        elif paragraph._p.xpath(".//w:drawing"):
            paragraph.style = document.styles["HIPHIL figure"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(14)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.keep_with_next = True
        elif bibliography_started:
            paragraph.style = document.styles["HIPHIL Bibliography Author-Date"]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.line_spacing = Pt(13)
            paragraph.paragraph_format.space_after = Pt(4)
            bibliography_paragraphs.append(paragraph)
        elif original_style.startswith("Footnote"):
            paragraph.style = document.styles["HIPHIL Footnote"]
        elif original_style.startswith("List Bullet"):
            paragraph.style = document.styles["HIPHIL bullet"]
        elif original_style.startswith("List Number"):
            paragraph.style = document.styles["HIPHIL Numbered list"]
        elif paragraph._p.xpath("./w:pPr/w:numPr"):
            paragraph.style = document.styles["HIPHIL bullet"]
        else:
            paragraph.style = document.styles["HIPHIL body"]
            paragraph.paragraph_format.space_after = Pt(7)

        if value.startswith("Figure "):
            paragraph.paragraph_format.keep_with_next = False
        elif value.startswith(("Example ", "Table ")):
            paragraph.paragraph_format.keep_with_next = True
        elif original_style != "Source Code" and index + 1 < len(document.paragraphs):
            next_value = document.paragraphs[index + 1].text.strip()
            if next_value.startswith("Figure "):
                paragraph.paragraph_format.keep_with_next = True

    # Author-date bibliographies are alphabetical, independent of citation order.
    if bibliography_paragraphs:
        parent = bibliography_paragraphs[0]._element.getparent()
        sorted_elements = [
            paragraph._element
            for paragraph in sorted(
                bibliography_paragraphs,
                key=lambda paragraph: paragraph.text.casefold().replace("ø", "o"),
            )
        ]
        for element in sorted_elements:
            parent.remove(element)
        section_properties = parent.find(qn("w:sectPr"))
        for element in sorted_elements:
            if section_properties is None:
                parent.append(element)
            else:
                parent.insert(parent.index(section_properties), element)

    table_alignments = {
        1: [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        2: [
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.LEFT,
            WD_ALIGN_PARAGRAPH.RIGHT,
        ],
        3: [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.RIGHT] * 4,
        4: [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.RIGHT] * 4,
        5: [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.RIGHT] * 4,
    }
    for table_index, table in enumerate(document.tables, 1):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        format_booktabs_table(table)
        for row_index, row in enumerate(table.rows):
            for column_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.style = document.styles[
                        "HIPHIL TableHeading" if row_index == 0 else "HIPHIL TableText"
                    ]
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.alignment = table_alignments[table_index][column_index]
                    for run in paragraph.runs:
                        run.font.name = "Times New Roman"
                        run.font.size = Pt(10.5)

        next_element = table._tbl.getnext()
        while next_element is not None and next_element.tag not in {
            qn("w:p"),
            qn("w:tbl"),
            qn("w:sectPr"),
        }:
            next_element = next_element.getnext()
        if next_element is not None and next_element.tag == qn("w:p"):
            following_paragraph = Paragraph(next_element, table._parent)
            following_paragraph.paragraph_format.space_before = Pt(
                16 if following_paragraph._p.xpath(".//w:drawing") else 12
            )

    # Template-compliant black hyperlinks, including those in footnotes.
    hyperlink_style = document.styles["Hyperlink"]
    hyperlink_style.font.color.rgb = RGBColor(0, 0, 0)
    for part in [document.part, *document.part.package.parts]:
        element = getattr(part, "element", None)
        if element is None:
            continue
        for run_properties in element.xpath(".//w:hyperlink//w:rPr"):
            color = run_properties.find(qn("w:color"))
            if color is None:
                color = OxmlElement("w:color")
                run_properties.append(color)
            color.set(qn("w:val"), "000000")

    # The journal assigns volume/issue metadata later; retain its header design without placeholders.
    for section in document.sections:
        for table in section.header.tables:
            if not table.rows or not table.rows[0].cells:
                continue
            paragraph = table.rows[0].cells[0].paragraphs[0]
            paragraph.clear()
            journal = paragraph.add_run("HIPHIL Novum")
            journal.italic = True
            paragraph.add_run("\t")
            paragraph.add_run("http://hiphil.org")

    document.save(DOCX)


def scrub_package() -> None:
    temporary = BUILD / "scrubbed.docx"
    with ZipFile(DOCX) as source, ZipFile(temporary, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.endswith(".xml") or item.filename.endswith(".rels"):
                data = re.sub(rb' w:author="[^"]*"', b' w:author=""', data)
                data = re.sub(rb' w:initials="[^"]*"', b' w:initials=""', data)
                data = re.sub(rb'<w:commentRangeStart[^>]*/>', b'', data)
                data = re.sub(rb'<w:commentRangeEnd[^>]*/>', b'', data)
                data = re.sub(rb'<w:commentReference[^>]*/>', b'', data)
                data = re.sub(rb'<w:ins[^>]*>(.*?)</w:ins>', rb'\1', data, flags=re.DOTALL)
                data = re.sub(rb'<w:del[^>]*>.*?</w:del>', b'', data, flags=re.DOTALL)
                if item.filename.endswith(".xml"):
                    root = etree.fromstring(data)
                    for changed_property in root.xpath("//*[contains(local-name(), 'PrChange')]"):
                        parent = changed_property.getparent()
                        if parent is not None:
                            parent.remove(changed_property)
                    data = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                if item.filename == "word/footnotes.xml":
                    data = re.sub(
                        rb'<w:pStyle w:val="Fodnotetekst"/>',
                        b'<w:pStyle w:val="HIPHILFootnote"/>',
                        data,
                    )
                if item.filename.startswith("word/comments") or item.filename.startswith("word/people"):
                    root = etree.fromstring(data)
                    for child in list(root):
                        root.remove(child)
                    data = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                elif item.filename == "docProps/custom.xml":
                    root = etree.fromstring(data)
                    for child in list(root):
                        root.remove(child)
                    data = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                elif item.filename == "docProps/app.xml":
                    root = etree.fromstring(data)
                    removable = {
                        "TotalTime",
                        "Pages",
                        "Words",
                        "Characters",
                        "Lines",
                        "Paragraphs",
                        "CharactersWithSpaces",
                        "HeadingPairs",
                        "TitlesOfParts",
                    }
                    for child in list(root):
                        if etree.QName(child).localname in removable:
                            root.remove(child)
                        elif etree.QName(child).localname == "Company":
                            child.text = ""
                    data = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
            target.writestr(item, data)
    shutil.move(temporary, DOCX)


def build_docx() -> None:
    ensure_template()
    render_assets()
    latex = prepare_latex()
    draft = BUILD / "pandoc.docx"
    run(
        PANDOC,
        "--from=latex",
        "--to=docx",
        "--standalone",
        "--reference-doc",
        TEMPLATE,
        "--resource-path",
        f"{ARTICLE_DIR}:{ASSETS}",
        latex,
        "--output",
        draft,
    )
    style_document(draft)
    scrub_package()


if __name__ == "__main__":
    build_docx()
    print(DOCX)
